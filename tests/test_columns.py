"""Multi-column sections: column geometry, flow order, and explicit column breaks.

`<w:cols w:num="2"/>` divides a section's text area into columns that content fills one at a time,
top to bottom, before moving right. None of the sample documents in the repo exercise this (every
`w:sectPr` in them is single-column), so these tests build their own fixtures.

Unlike `test_word_pagination.py`, there is no Word-recorded layout to check against here -- a
`.docx` built by python-docx carries no `<Pages>` and no `PAGEREF` results, because nothing has laid
it out. What *is* checkable without Word is the part the spec fixes exactly: column widths and
offsets are arithmetic, and the flow order is unambiguous. So these assert geometry and reading
order rather than pagination.
"""

from __future__ import annotations

from pathlib import Path

import docx
import pytest
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn as docx_qn
from docx.shared import Inches, Pt

from docx_renderer import DocxParser, FontManager, LayoutEngine
from docx_renderer.model import SectionProps

# A letter page with 1" margins: 612 - 72 - 72 = 468 pt of text width.
CONTENT_W = 468.0
#: Distinct words, so a reading-order comparison pinpoints where a re-break went wrong.
WORDS = "alpha bravo charlie delta echo foxtrot golf hotel india "
MARGIN_L = 72.0
MARGIN_TOP = 72.0


def _set_columns(section, num: int, space_twips: int = 720) -> None:
    """Set `<w:cols w:num=.. w:space=..>` on a section, reusing the element if it is already there."""
    sectPr = section._sectPr
    cols = sectPr.find(docx_qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sectPr.append(cols)
    cols.set(docx_qn("w:num"), str(num))
    cols.set(docx_qn("w:space"), str(space_twips))


def _add_column_break(paragraph) -> None:
    run = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(docx_qn("w:type"), "column")
    run.append(br)
    paragraph._p.append(run)


def _build(path: Path, n_paras: int = 60, cols: int = 2, break_before: int | None = None) -> Path:
    """A one-column heading, then a continuous `cols`-column section of numbered paragraphs."""
    doc = docx.Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(9)

    s0 = doc.sections[0]
    s0.top_margin = s0.bottom_margin = Inches(1.0)
    s0.left_margin = s0.right_margin = Inches(1.0)
    _set_columns(s0, 1)

    head = doc.add_paragraph("SINGLE COLUMN HEADING")
    head.paragraph_format.space_after = Pt(0)

    s1 = doc.add_section(WD_SECTION.CONTINUOUS)
    s1.top_margin = s1.bottom_margin = Inches(1.0)
    s1.left_margin = s1.right_margin = Inches(1.0)
    _set_columns(s1, cols)

    for i in range(1, n_paras + 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        if break_before is not None and i == break_before:
            _add_column_break(p)
        p.add_run(f"P{i:03d}")

    doc.save(str(path))
    return path


def _layout(path: Path):
    return LayoutEngine(DocxParser(path).parse(), font_manager=FontManager()).layout_document()


def _labelled_texts(pages):
    """Every drawn 'Pnnn' label as (page_number, x, y_baseline, text)."""
    out = []
    for pg in pages:
        for t in pg.texts:
            txt = (t.text or "").strip()
            if txt.startswith("P") and txt[1:].isdigit():
                out.append((pg.page_number, round(t.x, 2), t.y_baseline, txt))
    return out


# --------------------------------------------------------------------------------------------
# geometry -- pure arithmetic, no document needed
# --------------------------------------------------------------------------------------------

def test_single_column_spans_the_whole_text_area() -> None:
    sp = SectionProps()
    assert sp.column_bounds() == [(MARGIN_L, CONTENT_W)]


def test_two_equal_columns_split_the_text_area_minus_the_gutter() -> None:
    sp = SectionProps(cols_num=2, cols_space_pt=36.0)
    (x0, w0), (x1, w1) = sp.column_bounds()
    assert w0 == w1 == (CONTENT_W - 36.0) / 2.0 == 216.0
    assert x0 == MARGIN_L
    assert x1 == MARGIN_L + 216.0 + 36.0 == 324.0
    # the last column must end exactly on the right margin
    assert x1 + w1 == MARGIN_L + CONTENT_W == 540.0


def test_three_equal_columns() -> None:
    sp = SectionProps(cols_num=3, cols_space_pt=18.0)
    bounds = sp.column_bounds()
    assert len(bounds) == 3
    assert all(abs(w - (CONTENT_W - 2 * 18.0) / 3.0) < 1e-9 for _, w in bounds)
    assert abs(bounds[-1][0] + bounds[-1][1] - (MARGIN_L + CONTENT_W)) < 1e-9


def test_explicit_unequal_columns_follow_their_declared_widths() -> None:
    sp = SectionProps(
        cols_num=2,
        cols_equal_width=False,
        cols_explicit=[(300.0, 24.0), (144.0, 0.0)],
    )
    assert sp.column_bounds() == [(MARGIN_L, 300.0), (MARGIN_L + 324.0, 144.0)]


# --------------------------------------------------------------------------------------------
# parsing
# --------------------------------------------------------------------------------------------

def test_parser_reads_column_count_and_gutter(tmp_path: Path) -> None:
    doc = DocxParser(_build(tmp_path / "cols.docx", n_paras=4, cols=2)).parse()
    sects = [b.sect_pr for b in doc.blocks if getattr(b, "sect_pr", None) is not None]
    assert sects, "expected the continuous section break to be parsed"
    assert sects[0].cols_num == 1, "the first section is single-column"
    assert doc.final_sect_pr.cols_num == 2
    assert doc.final_sect_pr.cols_space_pt == pytest.approx(36.0)


# --------------------------------------------------------------------------------------------
# flow
# --------------------------------------------------------------------------------------------

def test_two_column_flow_fills_each_column_before_moving_right(tmp_path: Path) -> None:
    pages = _layout(_build(tmp_path / "two_col.docx", n_paras=60, cols=2))
    drawn = _labelled_texts(pages)
    assert len(drawn) == 60, "every paragraph should be drawn exactly once"

    xs = sorted({x for _, x, _, _ in drawn})
    assert xs == [MARGIN_L, 324.0], f"expected text at both column origins, got {xs}"

    # Reading order: page by page, left column top-to-bottom, then right column.
    order = [
        txt
        for pg in sorted({p for p, _, _, _ in drawn})
        for x in xs
        for _, _, _, txt in sorted(
            [d for d in drawn if d[0] == pg and d[1] == x], key=lambda d: d[2]
        )
    ]
    assert order == [f"P{i:03d}" for i in range(1, 61)]


def test_second_column_starts_level_with_the_column_set_not_the_page(tmp_path: Path) -> None:
    """A continuous break opens the column set part-way down; column 2 starts level with it.

    Resetting to the top margin instead would float the right column above the heading that the
    one-column section before it put on the same page.
    """
    pages = _layout(_build(tmp_path / "two_col.docx", n_paras=60, cols=2))
    drawn = _labelled_texts(pages)

    first_page = [d for d in drawn if d[0] == 1]
    left = [d for d in first_page if d[1] == MARGIN_L]
    right = [d for d in first_page if d[1] == 324.0]
    assert left and right, "page 1 should use both columns"

    # Both columns of the set begin on the same line, below the single-column heading above them.
    assert min(d[2] for d in right) == pytest.approx(min(d[2] for d in left), abs=0.01)
    assert min(d[2] for d in left) > MARGIN_TOP + 12.0, "the heading should sit above the set"
    assert max(d[2] for d in left) > 700.0, "the left column should fill to the bottom margin"
    # ...and the right column really is a continuation, starting above where the left one ended.
    assert min(d[2] for d in right) < max(d[2] for d in left)


def test_two_columns_need_fewer_pages_than_one(tmp_path: Path) -> None:
    one = _layout(_build(tmp_path / "one.docx", n_paras=120, cols=1))
    two = _layout(_build(tmp_path / "two.docx", n_paras=120, cols=2))
    assert len(two) < len(one)


def test_long_paragraph_wraps_inside_its_column(tmp_path: Path) -> None:
    """If the line breaker still measured the full text width, lines would overrun the column."""
    path = tmp_path / "wrap.docx"
    doc = docx.Document()
    doc.styles["Normal"].font.size = Pt(9)
    s = doc.sections[0]
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1.0)
    _set_columns(s, 2)
    doc.add_paragraph("word " * 400)
    doc.save(str(path))

    pages = _layout(path)
    assert any(pg.texts for pg in pages)
    for pg in pages:
        for t in pg.texts:
            # every line must sit inside one of the two columns
            assert (MARGIN_L - 0.5 <= t.x <= MARGIN_L + 216.0) or (324.0 - 0.5 <= t.x <= 540.0), (
                f"text at x={t.x:.1f} is outside both columns"
            )


def test_column_break_jumps_to_the_next_column(tmp_path: Path) -> None:
    """`<w:br w:type="column"/>` moves to the next column even though the current one has room."""
    pages = _layout(_build(tmp_path / "br.docx", n_paras=8, cols=2, break_before=5))
    drawn = _labelled_texts(pages)
    assert len(drawn) == 8

    before = {txt: x for _, x, _, txt in drawn if int(txt[1:]) < 5}
    after = {txt: x for _, x, _, txt in drawn if int(txt[1:]) >= 5}
    assert set(before.values()) == {MARGIN_L}, "P001-P004 belong in the left column"
    assert set(after.values()) == {324.0}, "P005 onward must be pushed into the right column"
    assert len(pages) == 1, "a column break must not start a new page while a column is free"


def _build_with_table(path: Path, rows: int) -> Path:
    """A 2-column section whose content is one table of `rows` rows."""
    doc = docx.Document()
    doc.styles["Normal"].font.size = Pt(9)
    s = doc.sections[0]
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1.0)
    _set_columns(s, 2)

    tbl = doc.add_table(rows=rows, cols=2)
    for r_i, row in enumerate(tbl.rows):
        for c_i, cell in enumerate(row.cells):
            cell.paragraphs[0].add_run(f"r{r_i:03d}c{c_i}")
    doc.save(str(path))
    return path


def test_table_in_a_column_is_laid_out_inside_that_column(tmp_path: Path) -> None:
    """A table in a two-column section belongs to its column, not to the full page width."""
    pages = _layout(_build_with_table(tmp_path / "tbl.docx", rows=6))
    boxes = [bb for pg in pages for bb in pg.bboxes if bb.kind == "table"]
    assert boxes, "expected a table bbox"
    for bb in boxes:
        width = bb.x1 - bb.x0
        assert width <= 216.0 + 0.5, f"table is {width:.1f}pt wide, wider than its 216pt column"
        in_left = MARGIN_L - 0.5 <= bb.x0 and bb.x1 <= MARGIN_L + 216.0 + 0.5
        in_right = 324.0 - 0.5 <= bb.x0 and bb.x1 <= 540.0 + 0.5
        assert in_left or in_right, f"table spans x={bb.x0:.1f}..{bb.x1:.1f}, outside both columns"


def test_long_table_continues_into_the_next_column(tmp_path: Path) -> None:
    """Rows that do not fit move to the next column before they start a new page."""
    pages = _layout(_build_with_table(tmp_path / "tall.docx", rows=90))
    xs = {round(t.x, 2) for pg in pages for t in pg.texts if (t.text or "").startswith("r")}
    assert any(x >= 324.0 for x in xs), "a table taller than one column must continue in column 2"
    assert len(pages) == 1, "it should fill both columns of page 1 before paginating"


def _build_balanced(path: Path, n_paras: int, trailing_section: bool) -> Path:
    """A 2-column section, optionally followed by a continuous single-column one.

    That trailing continuous break is what asks Word to balance the columns -- it is exactly what
    the "insert a continuous section break to balance columns" advice adds.
    """
    doc = docx.Document()
    doc.styles["Normal"].font.size = Pt(9)
    s0 = doc.sections[0]
    s0.top_margin = s0.bottom_margin = s0.left_margin = s0.right_margin = Inches(1.0)
    _set_columns(s0, 1)
    doc.add_paragraph("HEADING").paragraph_format.space_after = Pt(0)

    s1 = doc.add_section(WD_SECTION.CONTINUOUS)
    s1.top_margin = s1.bottom_margin = s1.left_margin = s1.right_margin = Inches(1.0)
    _set_columns(s1, 2)
    for i in range(1, n_paras + 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
        p.add_run(f"P{i:03d}")

    if trailing_section:
        s2 = doc.add_section(WD_SECTION.CONTINUOUS)
        s2.top_margin = s2.bottom_margin = s2.left_margin = s2.right_margin = Inches(1.0)
        _set_columns(s2, 1)
        doc.add_paragraph("AFTER THE COLUMNS")
    doc.save(str(path))
    return path


def test_short_section_is_not_balanced_without_a_trailing_continuous_break(tmp_path: Path) -> None:
    """Word fills to the bottom margin unless a continuous break follows; so do we."""
    pages = _layout(_build_balanced(tmp_path / "unbal.docx", 20, trailing_section=False))
    xs = {x for _, x, _, _ in _labelled_texts(pages)}
    assert xs == {MARGIN_L}, "without the trailing break all 20 fit in the first column"


def test_trailing_continuous_break_balances_the_columns(tmp_path: Path) -> None:
    pages = _layout(_build_balanced(tmp_path / "bal.docx", 20, trailing_section=True))
    drawn = _labelled_texts(pages)
    assert len(drawn) == 20

    left = [d for d in drawn if d[1] == MARGIN_L]
    right = [d for d in drawn if d[1] == 324.0]
    assert right, "the right column must be used once the section is balanced"
    # Even to within a line. The section's own break mark is an extra, invisible line in the
    # right column, so the visible counts differ by one more than the balance does.
    assert abs(len(left) - len(right)) <= 2, f"lopsided: {len(left)} vs {len(right)}"
    # ...and the columns must end at about the same height.
    assert max(d[2] for d in left) - max(d[2] for d in right) <= 2 * 12.5


def test_balancing_shortens_the_block_against_the_unbalanced_layout(tmp_path: Path) -> None:
    """The point of balancing: a short even block, not one column run to the bottom margin."""
    unbalanced = _layout(_build_balanced(tmp_path / "u.docx", 40, trailing_section=False))
    balanced = _layout(_build_balanced(tmp_path / "b.docx", 40, trailing_section=True))

    u_bottom = max(d[2] for d in _labelled_texts(unbalanced))
    b_bottom = max(d[2] for d in _labelled_texts(balanced))
    assert b_bottom < u_bottom - 100.0, (
        f"balanced content reaches y={b_bottom:.0f}, unbalanced y={u_bottom:.0f} -- "
        "balancing should have roughly halved the height"
    )


def test_section_after_a_column_set_clears_the_tallest_column(tmp_path: Path) -> None:
    """The cursor ends at the shortest column; resuming there would overlap the tallest one."""
    for n in (3, 4, 5, 20, 21):
        pages = _layout(_build_balanced(tmp_path / f"after{n}.docx", n, trailing_section=True))
        boxes = [bb for bb in pages[0].bboxes if bb.kind == "paragraph"]
        cols = [bb for bb in boxes if (bb.summary or "").startswith("P")]
        after = [bb for bb in boxes if "AFTER" in (bb.summary or "")]
        assert cols and after, f"n={n}: expected both the columns and the section after them"
        assert min(bb.y0 for bb in after) >= max(bb.y1 for bb in cols) - 0.01, (
            f"n={n}: the following section overlaps the bottom of the tallest column"
        )


def test_only_the_last_page_of_a_tall_section_is_balanced(tmp_path: Path) -> None:
    """Earlier pages fill to the bottom; only the final column set is evened up."""
    pages = _layout(_build_balanced(tmp_path / "tall.docx", 400, trailing_section=True))
    drawn = _labelled_texts(pages)
    assert len(drawn) == 400, "every paragraph placed exactly once"
    assert len({t for _, _, _, t in drawn}) == 400, "no paragraph duplicated by the rewind"
    assert len(pages) >= 3

    last = max(p for p, _, _, _ in drawn)
    on_last = [d for d in drawn if d[0] == last]
    left = [d for d in on_last if d[1] == MARGIN_L]
    right = [d for d in on_last if d[1] == 324.0]
    assert left and right, "the final page should use both columns"
    assert abs(len(left) - len(right)) <= 2, f"final page lopsided: {len(left)} vs {len(right)}"

    # An earlier page is filled, not balanced: its columns run to the bottom margin.
    earlier = [d for d in drawn if d[0] == last - 1]
    assert max(d[2] for d in earlier) > 650.0, "pages before the last should fill to the margin"


def test_balancing_works_when_the_final_set_began_mid_block(tmp_path: Path) -> None:
    """The block that crosses onto the last page is checkpointed on the page before it.

    Requiring a rewind point already sitting at the top of the final column set silently skipped
    balancing for every section longer than one page.
    """
    pages = _layout(_build_balanced(tmp_path / "multi.docx", 160, trailing_section=True))
    drawn = _labelled_texts(pages)
    last = max(p for p, _, _, _ in drawn)
    assert last > 1, "this fixture must span more than one page"
    left = [d for d in drawn if d[0] == last and d[1] == MARGIN_L]
    right = [d for d in drawn if d[0] == last and d[1] == 324.0]
    assert right, "the last page's right column must be used"
    assert abs(len(left) - len(right)) <= 2, f"not balanced: {len(left)} vs {len(right)}"


def _set_explicit_columns(section, specs) -> None:
    """Declare unequal columns: `specs` is a list of (width_twips, space_twips)."""
    sectPr = section._sectPr
    cols = sectPr.find(docx_qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sectPr.append(cols)
    cols.set(docx_qn("w:equalWidth"), "0")
    cols.set(docx_qn("w:num"), str(len(specs)))
    for width, space in specs:
        col = OxmlElement("w:col")
        col.set(docx_qn("w:w"), str(width))
        col.set(docx_qn("w:space"), str(space))
        cols.append(col)


def _build_unequal(path: Path, reps: int = 120, first_line_indent_pt: float = 0.0) -> Path:
    """One paragraph long enough to flow from a 300pt column into a 144pt one."""
    doc = docx.Document()
    doc.styles["Normal"].font.size = Pt(9)
    s = doc.sections[0]
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1.0)
    _set_explicit_columns(s, [(6000, 480), (2880, 0)])   # 300pt + 24pt gutter + 144pt = 468pt
    p = doc.add_paragraph(WORDS * reps)
    if first_line_indent_pt:
        p.paragraph_format.first_line_indent = Pt(first_line_indent_pt)
    doc.save(str(path))
    return path


def _column_of(x: float, col_xs) -> float:
    """The origin of the column a drawn run sits in."""
    at_or_left = [c for c in col_xs if x >= c - 0.5]
    return max(at_or_left) if at_or_left else col_xs[0]


def _reading_order_words(pages, col_xs):
    """Every drawn word, in reading order: by page, then column, then line, then x."""
    words = []
    for pg in sorted(pages, key=lambda p: p.page_number):
        for cx in col_xs:
            lines = {}
            for t in pg.texts:
                if _column_of(t.x, col_xs) == cx:
                    lines.setdefault(round(t.y_baseline, 2), []).append(t)
            for y in sorted(lines):
                for t in sorted(lines[y], key=lambda run: run.x):
                    words.extend((t.text or "").split())
    return words


def test_unequal_columns_wrap_at_their_own_width(tmp_path: Path) -> None:
    """A paragraph flowing into a narrower column must be re-broken against it.

    Keeping the wide column's measure runs the narrow column's lines clean off the page.
    """
    pages = _layout(_build_unequal(tmp_path / "uneq.docx"))
    narrow_x = 396.0                      # 72 + 300 + 24
    right_edge = narrow_x + 144.0         # == 540, the page's right text edge

    in_narrow = [t for pg in pages for t in pg.texts if t.x >= narrow_x - 0.5]
    assert in_narrow, "the paragraph should have flowed into the second column"
    assert max(t.x for t in in_narrow) <= right_edge, (
        f"text reaches x={max(t.x for t in in_narrow):.1f}, past the {right_edge:.0f}pt right edge"
    )


def test_reflowing_into_a_narrower_column_loses_no_words(tmp_path: Path) -> None:
    """The seam is where a re-break could duplicate or drop content."""
    reps = 120
    pages = _layout(_build_unequal(tmp_path / "uneq.docx", reps=reps))
    got = _reading_order_words(pages, [MARGIN_L, 396.0])
    expected = (WORDS * reps).split()
    assert got == expected, (
        f"{len(got)} words laid out, expected {len(expected)}; "
        f"first difference at {next((i for i, (a, b) in enumerate(zip(got, expected)) if a != b), 'none')}"
    )


def test_continuation_in_a_new_column_is_not_first_line_indented(tmp_path: Path) -> None:
    """Only the line the paragraph starts on carries the first-line indent."""
    first_page = _layout(_build_unequal(tmp_path / "ind.docx", first_line_indent_pt=36.0))[0]

    def first_line_x(runs):
        top = min(t.y_baseline for t in runs)
        return min(t.x for t in runs if abs(t.y_baseline - top) < 0.01)

    narrow = [t for t in first_page.texts if t.x >= 396.0 - 0.5]
    wide = [t for t in first_page.texts if t.x < 396.0 - 0.5]
    assert narrow and wide, "expected the paragraph to span both columns"

    # The paragraph's own first line is indented...
    assert first_line_x(wide) == pytest.approx(MARGIN_L + 36.0, abs=0.5)
    # ...the continuation at the top of the next column is not.
    assert first_line_x(narrow) == pytest.approx(396.0, abs=0.5), (
        "the continuation should start at the column edge, not indented like a new paragraph"
    )


def test_single_column_document_is_unchanged_by_column_support(tmp_path: Path) -> None:
    pages = _layout(_build(tmp_path / "one_col.docx", n_paras=20, cols=1))
    for _, x, _, _ in _labelled_texts(pages):
        assert x == MARGIN_L
