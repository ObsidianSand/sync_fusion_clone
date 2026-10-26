"""End-to-end tests and sample document generator for the docx_renderer pipeline."""

from __future__ import annotations
import io
import json
from pathlib import Path
from PIL import Image, ImageDraw
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn as docx_qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from docx_renderer import render_docx_to_png, DocxParser, LayoutEngine


def _make_sample_figure_png() -> bytes:
    """Create a synthetic architecture figure PNG in memory for embedding into the test .docx."""
    img = Image.new("RGB", (640, 220), color=(246, 249, 253))
    draw = ImageDraw.Draw(img)
    # Border
    draw.rectangle([2, 2, 637, 217], outline=(68, 114, 196), width=2)
    # 3 pipeline boxes
    boxes = [
        (40, 60, 190, 155, (68, 114, 196), "1. OpenXML\n   Parser"),
        (245, 60, 395, 155, (47, 85, 151), "2. Box/Flow\n   Layout"),
        (450, 60, 600, 155, (84, 130, 53), "3. PyMuPDF\n   PNG Render"),
    ]
    for x0, y0, x1, y1, col, label in boxes:
        draw.rectangle([x0, y0, x1, y1], fill=col, outline=(30, 50, 90), width=2)
        draw.text((x0 + 25, y0 + 32), label, fill=(255, 255, 255))
    # Connecting arrows
    for ax0, ax1 in [(190, 245), (395, 450)]:
        draw.line([(ax0, 107), (ax1, 107)], fill=(68, 114, 196), width=4)
        draw.polygon([(ax1, 107), (ax1 - 10, 100), (ax1 - 10, 114)], fill=(68, 114, 196))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _add_toc_entry(body_or_sdt, title: str, page_str: str, level: int = 1) -> None:
    """Add a realistic OpenXML TOC paragraph with right-aligned dot-leader tab stop at 6.5 inches (9360 twips)."""
    p = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")

    pStyle = OxmlElement("w:pStyle")
    pStyle.set(docx_qn("w:val"), f"TOC{level}")
    pPr.append(pStyle)

    # Indentation per TOC level (0", 0.25", 0.50")
    ind = OxmlElement("w:ind")
    ind.set(docx_qn("w:left"), str((level - 1) * 360))
    pPr.append(ind)

    # Spacing
    spc = OxmlElement("w:spacing")
    spc.set(docx_qn("w:before"), "40" if level == 1 else "0")
    spc.set(docx_qn("w:after"), "60")
    pPr.append(spc)

    # Right-aligned tab stop with dot leader at 9360 twips (468 pt = 6.5 in)
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(docx_qn("w:val"), "right")
    tab.set(docx_qn("w:leader"), "dot")
    tab.set(docx_qn("w:pos"), "9360")
    tabs.append(tab)
    pPr.append(tabs)
    p.append(pPr)

    # Title run
    r1 = OxmlElement("w:r")
    if level == 1:
        rPr1 = OxmlElement("w:rPr")
        b = OxmlElement("w:b")
        rPr1.append(b)
        r1.append(rPr1)
    t1 = OxmlElement("w:t")
    t1.text = title
    r1.append(t1)
    p.append(r1)

    # Tab run
    rtab = OxmlElement("w:r")
    rtab.append(OxmlElement("w:tab"))
    p.append(rtab)

    # Page number run
    r2 = OxmlElement("w:r")
    t2 = OxmlElement("w:t")
    t2.text = page_str
    r2.append(t2)
    p.append(r2)

    body_or_sdt.append(p)


def _set_cell_shading(cell, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(docx_qn("w:val"), "clear")
    shd.set(docx_qn("w:color"), "auto")
    shd.set(docx_qn("w:fill"), hex_color)
    tcPr.append(shd)


def _set_table_borders(table, color: str = "BFBFBF", sz: str = "6") -> None:
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "bottom", "left", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{side}")
        b.set(docx_qn("w:val"), "single")
        b.set(docx_qn("w:sz"), sz)
        b.set(docx_qn("w:space"), "0")
        b.set(docx_qn("w:color"), color)
        borders.append(b)
    tblPr.append(borders)


def create_showcase_docx(output_path: str | Path) -> Path:
    """Create a comprehensive Word document showcasing TOC, Indents/Tabs, Tables, and Figures."""
    doc = docx.Document()

    # Set modern Word default font (Calibri) on Normal style
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(10)

    # Configure header & footer
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    hdr_p = section.header.paragraphs[0]
    hdr_p.text = "SEMOSS MCP — High-Fidelity OpenXML Layout & Rendering Showcase"
    hdr_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hdr_p.runs[0].font.name = "Calibri"
    hdr_p.runs[0].font.size = Pt(8.5)
    hdr_p.runs[0].font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    ftr_p = section.footer.paragraphs[0]
    ftr_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ftr = ftr_p.add_run("Confidential — Page ")
    r_ftr.font.name = "Calibri"
    r_ftr.font.size = Pt(9)
    # Add PAGE field
    fldSimple = OxmlElement("w:fldSimple")
    fldSimple.set(docx_qn("w:instr"), "PAGE")
    ftr_p._p.append(fldSimple)

    # Document Title
    title_p = doc.add_heading("Technical Specification: OpenXML Visual Feedback Pipeline", level=0)
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_p.paragraph_format.space_after = Pt(4)
    for r in title_p.runs:
        r.font.size = Pt(18)

    # Table of Contents inside an authentic <w:sdt> block
    toc_heading = doc.add_heading("Table of Contents", level=1)
    toc_heading.paragraph_format.space_before = Pt(4)
    toc_heading.paragraph_format.space_after = Pt(2)

    sdt = OxmlElement("w:sdt")
    sdtContent = OxmlElement("w:sdtContent")
    _add_toc_entry(sdtContent, "1. Executive Summary & Indentation Verification", "1", level=1)
    _add_toc_entry(sdtContent, "1.1 Tab Stop & Hanging Indent Alignment", "1", level=2)
    _add_toc_entry(sdtContent, "1.2 Multi-Level Lists & Sub-Clauses", "1", level=2)
    _add_toc_entry(sdtContent, "2. Performance & Financial Metrics Table", "1", level=1)
    _add_toc_entry(sdtContent, "3. System Architecture Figure", "1", level=1)
    sdt.append(sdtContent)
    toc_heading._p.addnext(sdt)

    # Section 1: Indentation & Tab Alignment Showcase
    h1 = doc.add_heading("1. Executive Summary & Indentation Verification", level=1)
    h1.paragraph_format.space_before = Pt(8)
    h1.paragraph_format.space_after = Pt(2)

    p_intro = doc.add_paragraph(
        "This document tests exact spatial positioning in Twips and EMUs. When an LLM edits raw OpenXML, "
        "it relies on rendered page PNGs to verify paragraph indents, tab stops, table column widths, and figures."
    )
    p_intro.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_intro.paragraph_format.space_after = Pt(4)

    h2 = doc.add_heading("1.1 Tab Stop & Hanging Indent Alignment", level=2)
    h2.paragraph_format.space_before = Pt(6)
    h2.paragraph_format.space_after = Pt(2)

    # Demonstrate multi-level clause indentation (including one intentionally shifted by 1 tab so the LLM can spot it!)
    clauses = [
        (0.0, "Clause 1.0 (0.0\" margin): Properly aligned at the primary left margin."),
        (0.5, "Clause 1.1 (0.5\" / 1 Tab): Indented exactly one standard tab stop over."),
        (1.0, "Clause 1.1.1 (1.0\" / 2 Tabs): Indented two tab stops for nested sub-items."),
        (1.5, "Clause 1.1.2 (1.5\" / 3 Tabs — MISMATCH DEMO): Intentionally indented one extra tab over!"),
        (1.0, "Clause 1.1.3 (1.0\" / 2 Tabs): Returned to the expected two-tab sibling alignment."),
    ]
    for indent_in, text in clauses:
        cp = doc.add_paragraph()
        cp.paragraph_format.left_indent = Inches(indent_in)
        cp.paragraph_format.space_before = Pt(0)
        cp.paragraph_format.space_after = Pt(2)
        run = cp.add_run(text)
        run.font.size = Pt(9.5)
        if "MISMATCH" in text:
            run.bold = True
            run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    # Section 1.2: Multi-Level Lists
    h2_list = doc.add_heading("1.2 Multi-Level Lists & Sub-Clauses", level=2)
    h2_list.paragraph_format.space_before = Pt(6)
    h2_list.paragraph_format.space_after = Pt(2)

    for b_text in [
        "Primary bullet item rendered via numbering.xml with hanging indent.",
        "Secondary bullet item verifying consistent horizontal alignment after bullet marker.",
    ]:
        bp = doc.add_paragraph(b_text, style="List Bullet")
        bp.paragraph_format.space_before = Pt(0)
        bp.paragraph_format.space_after = Pt(2)

    # Section 2: Complex Table with Merged Cells, Shading & Alignment
    h_tbl = doc.add_heading("2. Performance & Financial Metrics Table", level=1)
    h_tbl.paragraph_format.space_before = Pt(8)
    h_tbl.paragraph_format.space_after = Pt(3)

    tbl = doc.add_table(rows=5, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(tbl, color="8EA9DB", sz="6")

    col_widths = [Inches(1.6), Inches(1.9), Inches(1.5), Inches(1.5)]
    headers = ["Component", "Subsystem / Metric", "Target SLA", "Measured"]
    hdr_row = tbl.rows[0]
    for c_idx, h_text in enumerate(headers):
        cell = hdr_row.cells[c_idx]
        cell.width = col_widths[c_idx]
        _set_cell_shading(cell, "2F5496")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx >= 2 else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(h_text)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    table_data = [
        ("Layout Engine", "Twips / Tab Stop Precision", "±0.5 pt", "±0.1 pt"),
        ("", "Table Grid & Merge Resolution", "100%", "100%"),
        ("Canvas Renderer", "150 DPI Page Export Latency", "< 50 ms", "18 ms"),
    ]
    for r_offset, row_vals in enumerate(table_data, start=1):
        row = tbl.rows[r_offset]
        bg = "D9E2F3" if r_offset % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_vals):
            cell = row.cells[c_idx]
            cell.width = col_widths[c_idx]
            _set_cell_shading(cell, bg)
            if val:
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if c_idx >= 2 else WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(val)
                r.font.size = Pt(9.5)
                if c_idx == 0 and r_offset == 1:
                    r.bold = True

    # Merge "Layout Engine" vertically across rows 1 and 2 and center vertically
    merged_v = tbl.rows[1].cells[0].merge(tbl.rows[2].cells[0])
    vAlign = OxmlElement("w:vAlign")
    vAlign.set(docx_qn("w:val"), "center")
    merged_v._tc.get_or_add_tcPr().append(vAlign)

    # Summary row with horizontal span across first 2 columns
    sum_row = tbl.rows[4]
    sum_row.cells[0].merge(sum_row.cells[1])
    for c_idx in range(4):
        _set_cell_shading(sum_row.cells[c_idx], "EDEDED")
    p_sum = sum_row.cells[0].paragraphs[0]
    r_sum = p_sum.add_run("Overall Pipeline Readiness (Offline Linux / SEMOSS MCP)")
    r_sum.bold = True
    r_sum.font.size = Pt(9.5)
    p_stat = sum_row.cells[2].paragraphs[0]
    p_stat.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_stat.add_run("99.9%").bold = True
    p_pass = sum_row.cells[3].paragraphs[0]
    p_pass.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_pass = p_pass.add_run("PASSED")
    r_pass.bold = True
    r_pass.font.color.rgb = RGBColor(0x38, 0x57, 0x23)

    # Section 3: Embedded Figure & Caption
    h_fig = doc.add_heading("3. System Architecture Figure", level=1)
    h_fig.paragraph_format.space_before = Pt(8)
    h_fig.paragraph_format.space_after = Pt(3)

    fig_bytes = _make_sample_figure_png()
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_after = Pt(2)
    p_img.add_run().add_picture(io.BytesIO(fig_bytes), width=Inches(4.3))

    p_cap = doc.add_paragraph("Figure 1: Three-stage offline OpenXML to PNG rendering architecture.")
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.runs[0].italic = True
    p_cap.runs[0].font.size = Pt(9)
    p_cap.runs[0].font.color.rgb = RGBColor(0x44, 0x54, 0x6A)

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


def test_full_render_pipeline(tmp_path: Path) -> None:
    docx_file = create_showcase_docx(tmp_path / "showcase.docx")
    assert docx_file.exists()

    # 1. Standard High-Fidelity Render
    clean_pngs = render_docx_to_png(
        docx_file,
        output_dir=tmp_path / "output_clean",
        prefix="clean_page",
        dpi=150,
        show_ruler=False,
        show_bboxes=False,
        export_layout_json=True,
    )
    assert len(clean_pngs) >= 1
    assert clean_pngs[0].exists()

    # Verify image dimensions at 150 DPI for 8.5 x 11 inch page (1275 x 1650 px)
    with Image.open(clean_pngs[0]) as im:
        assert abs(im.width - 1275) <= 2
        assert abs(im.height - 1650) <= 2

    # 2. Verify Layout JSON contains exact tab-indent measurements for the clauses
    layout_json = json.loads((tmp_path / "output_clean" / "clean_page_layout.json").read_text(encoding="utf-8"))
    elements = layout_json[0]["elements"]
    mismatch_elems = [e for e in elements if "MISMATCH DEMO" in e.get("summary", "")]
    assert len(mismatch_elems) == 1
    # Clause 1.1.2 was indented 1.5 inches = 3.0 tabs (108 pt from left margin)
    assert abs(mismatch_elems[0]["indent_in_tabs_0_5in"] - 3.0) < 0.05
    assert abs(mismatch_elems[0]["indent_from_left_margin_pt"] - 108.0) < 1.0

    # Verify TOC dot leaders were emitted
    toc_leaders = [e for e in elements if e["kind"] == "tab_leader"]
    assert len(toc_leaders) == 5

    # 3. Render with LLM Visual Ruler & Bounding Boxes
    llm_pngs = render_docx_to_png(
        docx_file,
        output_dir=tmp_path / "output_llm_debug",
        prefix="llm_page",
        dpi=150,
        show_ruler=True,
        show_bboxes=True,
    )
    assert len(llm_pngs) >= 1
    assert llm_pngs[0].exists()
